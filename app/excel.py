"""Génération de fichiers Excel (.xlsx) pour les exports."""

from collections.abc import Iterable, Sequence
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COLUMN_WIDTH = 40


def build_xlsx(sheet_title: str, headers: Sequence[str], rows: Iterable[Sequence]) -> bytes:
    """Classeur d'une feuille : en-tête en gras et figé, filtre automatique, colonnes ajustées au contenu.

    Les date et datetime deviennent de vraies dates Excel, les booléens Yes / No, le reste est écrit tel quel
    (les chaînes restent du texte : pas de MSISDN ou de numéro d'identité convertis en nombre).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    widths = [len(header) for header in headers]
    for row in rows:
        values = [_cell_value(value) for value in row]
        sheet.append(values)
        for index, value in enumerate(values):
            widths[index] = max(widths[index], len(_display(value)))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_value(value):
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime) and value.tzinfo is not None:
        # Excel ne gère pas les fuseaux : heure locale du serveur, comme à l'écran
        return value.astimezone().replace(tzinfo=None)
    return value


def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value:%Y-%m-%d %H:%M:%S}"
    return str(value)
